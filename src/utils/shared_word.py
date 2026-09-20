"""Editable Word tables for manuscript exports."""

from pathlib import Path


def save_word_table(frame, path, *, title, notes=(), column_weights=None):
    """Write a labelled DataFrame index and columns as a landscape A4 Word table.

    Callers supply display-formatted values; the source numerical table remains
    available separately as CSV. Word retains native editable cells and a repeated
    header, rather than an embedded image of the table.
    """
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    display = frame.reset_index()
    weights = column_weights or [2] + [1] * (len(display.columns) - 1)
    if len(weights) != len(display.columns) or any(value <= 0 for value in weights):
        raise ValueError("Supply one positive column weight for each displayed column")

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Mm(297), Mm(210)
    section.left_margin = section.right_margin = Mm(12)
    section.top_margin = section.bottom_margin = Mm(12)
    normal = document.styles["Normal"]
    normal.font.name = "Helvetica"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(7)
    run = heading.add_run(title)
    run.bold, run.font.size = True, Pt(13)

    table = document.add_table(rows=1, cols=len(display.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = section.page_width - section.left_margin - section.right_margin
    widths = [int(width * value / sum(weights)) for value in weights]
    for column, cell_width in zip(table.columns, widths):
        column.width = cell_width
    for cell, label in zip(table.rows[0].cells, display.columns):
        cell.text = str(label)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F2F2")
        cell._tc.get_or_add_tcPr().append(shading)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for values in display.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.text = str(value)
            if index:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, cell_width in zip(row.cells, widths):
            cell.width = cell_width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)

    for index, note in enumerate(notes):
        paragraph = document.add_paragraph(str(note))
        paragraph.paragraph_format.space_before = Pt(6 if index == 0 else 2)
        for run in paragraph.runs:
            run.font.size = Pt(8)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path
